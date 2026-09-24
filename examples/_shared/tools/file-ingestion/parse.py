"""Application-owned parsers. Prints JSON only; never receives expected test answers."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import subprocess

parser = argparse.ArgumentParser()
parser.add_argument("kind", choices=["pdf", "spreadsheet", "image", "audio"])
parser.add_argument("path")
parser.add_argument("--media-type", required=True)
parser.add_argument("--pdftotext", default="pdftotext")
parser.add_argument("--tesseract", default="tesseract")
parser.add_argument("--asr-model", default="tiny.en")
parser.add_argument("--asr-language", default="en")
args = parser.parse_args()
path = Path(args.path)
data = path.read_bytes()
result = {"name": path.name, "mediaType": args.media_type, "sourceSha256": hashlib.sha256(data).hexdigest()}
if args.kind == "pdf":
    if args.media_type != "application/pdf" or not data.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF signature")
    result["text"] = subprocess.run([args.pdftotext, "-layout", str(path), "-"], check=True, capture_output=True, text=True, timeout=60).stdout.strip()
    result["engine"] = "Poppler pdftotext"
elif args.kind == "spreadsheet":
    if args.media_type == "text/csv":
        result["rows"] = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True))
        result["engine"] = "Python csv"
    elif args.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            result["rows"] = [["" if value is None else str(value) for value in row] for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()
        result["engine"] = "openpyxl"
    else:
        raise ValueError("Unsupported spreadsheet MIME")
elif args.kind == "image":
    from PIL import Image
    with Image.open(path) as image:
        expected = {"image/png": "PNG", "image/jpeg": "JPEG"}.get(args.media_type)
        if expected is None or image.format != expected:
            raise ValueError("Image bytes do not match MIME")
        image.verify()
    result["ocrText"] = subprocess.run([args.tesseract, str(path), "stdout", "-l", "eng", "--psm", "6"], check=True, capture_output=True, text=True, timeout=60).stdout.strip()
    result["engine"] = "Tesseract eng"
else:
    if args.media_type == "audio/wav":
        if data[:4] != b"RIFF" or data[8:12] != b"WAVE":
            raise ValueError("Invalid WAV signature")
    elif args.media_type == "audio/mpeg":
        if not (data[:3] == b"ID3" or (len(data) > 1 and data[0] == 255 and data[1] & 224 == 224)):
            raise ValueError("Invalid MP3 signature")
    else:
        raise ValueError("Unsupported audio MIME")
    from faster_whisper import WhisperModel
    model = WhisperModel(args.asr_model, device="cpu", compute_type="int8")
    segments, info = model.transcribe(str(path), language=None if args.asr_language == "auto" else args.asr_language, beam_size=5, condition_on_previous_text=False)
    result["segments"] = [{"start": round(segment.start, 3), "end": round(segment.end, 3), "text": segment.text.strip()} for segment in segments]
    result["transcript"] = " ".join(segment["text"] for segment in result["segments"]).strip()
    result["engine"] = "faster-whisper/" + args.asr_model
if args.kind != "spreadsheet" and not result.get("text", result.get("ocrText", result.get("transcript", ""))):
    raise ValueError("File parser returned no content")
print(json.dumps(result))
